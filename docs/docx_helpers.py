import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os
import sys

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:left w:val="none"/>
                <w:right w:val="none"/>
                <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:insideV w:val="none"/>
            </w:tblBorders>
        ''')
        tblPr[0].append(borders)

def format_table(table, header_bg="4F46E5", alt_bg="F8FAFC", border_color="E2E8F0"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color=border_color)
    for i, row in enumerate(table.rows):
        is_header = (i == 0)
        trPr = row._element.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if is_header:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
            if is_header:
                set_cell_background(cell, header_bg)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.size = Pt(9.5)
                        run.font.name = 'Calibri'
            else:
                if i % 2 == 1:
                    set_cell_background(cell, "FFFFFF")
                else:
                    set_cell_background(cell, alt_bg)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    for run in p.runs:
                        run.font.size = Pt(9)
                        run.font.name = 'Calibri'
                        run.font.color.rgb = RGBColor(30, 41, 59)

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    run = h.runs[0]
    run.font.name = 'Segoe UI'
    if level == 1:
        h.paragraph_format.space_before = Pt(20)
        h.paragraph_format.space_after = Pt(8)
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 41, 59) # Slate-800
        # Add bottom border under Heading 1
        pPr = h._element.get_or_add_pPr()
        pBdr = parse_xml(f'''
            <w:pBdr {nsdecls("w")}>
                <w:bottom w:val="single" w:sz="12" w:space="4" w:color="6366F1"/>
            </w:pBdr>
        ''')
        pPr.append(pBdr)
    elif level == 2:
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(79, 70, 229) # Indigo-600
    elif level == 3:
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(14, 116, 144) # Cyan-700
    elif level == 4:
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(3)
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(71, 85, 105) # Slate-600
    return h

def add_callout(doc, text, title="CATATAN ARSITEKTUR", color="4F46E5", bg="F5F3FF"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    r_title = p.add_run(f"📌 {title}: ")
    r_title.bold = True
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(9.5)
    r_title.font.color.rgb = RGBColor(79, 70, 229)
    
    r_text = p.add_run(text)
    r_text.font.name = 'Calibri'
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = RGBColor(30, 41, 59)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_text, lang=""):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "0F172A") # Slate-900 dark background
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="334155"/>
            <w:left w:val="single" w:sz="6" w:space="0" w:color="334155"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="334155"/>
            <w:right w:val="single" w:sz="6" w:space="0" w:color="334155"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(code_text.strip())
    r.font.name = 'Consolas'
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(241, 245, 249) # Slate-100
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

print("Helper functions defined successfully.")
